"""RBT12 do Simples Nacional: janela móvel de 12 meses, importação e carga da folha."""

import csv
import io
import re

from tributacao import normalizar_regime_apuracao, parse_mes


MESES_PT = {
    "jan": 1, "janeiro": 1, "fev": 2, "fevereiro": 2, "mar": 3, "marco": 3, "março": 3,
    "abr": 4, "abril": 4, "mai": 5, "maio": 5, "jun": 6, "junho": 6,
    "jul": 7, "julho": 7, "ago": 8, "agosto": 8, "set": 9, "setembro": 9,
    "out": 10, "outubro": 10, "nov": 11, "novembro": 11, "dez": 12, "dezembro": 12,
}


def competencia_add(comp, meses):
    ano, mes = parse_mes(comp)
    total = ano * 12 + (mes - 1) + int(meses)
    return f"{total // 12:04d}-{(total % 12) + 1:02d}"


def janela_competencias(mes_apuracao):
    """Doze competências imediatamente anteriores ao mês de apuração (não inclui o mês vigente)."""
    return [competencia_add(mes_apuracao, -i) for i in range(12, 0, -1)]


def parse_moeda_livre(bruto):
    if bruto is None:
        return 0.0
    s = str(bruto).strip().replace("R$", "").replace("\xa0", "").replace(" ", "")
    if not s:
        return 0.0
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return 0.0


def parse_competencia(texto):
    if texto is None:
        return None
    if hasattr(texto, "strftime"):
        return texto.strftime("%Y-%m")
    s = str(texto).strip().lower()
    m = re.search(r"(20\d{2})[-/\.](\d{1,2})", s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}"
    m = re.search(r"(\d{1,2})[-/\.](20\d{2})", s)
    if m:
        return f"{m.group(2)}-{int(m.group(1)):02d}"
    m = re.search(r"([a-zç]+)\s*/?\s*(20\d{2})", s)
    if m and m.group(1) in MESES_PT:
        return f"{m.group(2)}-{MESES_PT[m.group(1)]:02d}"
    if re.fullmatch(r"20\d{2}-\d{2}", s):
        return s
    return None


def upsert_competencia(cursor, competencia, receita=None, folha=None, origem="manual", observacao=None):
    cursor.execute(
        """
        INSERT INTO simples_competencias (competencia, receita_bruta, folha_encargos, origem, observacao, atualizado_em)
        VALUES (%s, %s, %s, %s, %s, NOW())
        ON CONFLICT (competencia) DO UPDATE SET
            receita_bruta = COALESCE(%s, simples_competencias.receita_bruta),
            folha_encargos = COALESCE(%s, simples_competencias.folha_encargos),
            origem = %s,
            observacao = COALESCE(%s, simples_competencias.observacao),
            atualizado_em = NOW()
        """,
        (
            competencia,
            receita or 0,
            folha or 0,
            origem,
            observacao,
            receita,
            folha,
            origem,
            observacao,
        ),
    )


_ORIGENS_CONGELADAS = {"manual", "planilha", "pgdas"}


def primeira_competencia_sistema(cursor, regime_apuracao="competencia"):
    cursor.execute("SELECT MIN(competencia) AS c FROM simples_competencias")
    row = cursor.fetchone() or {}
    if row.get("c"):
        return row["c"]
    if normalizar_regime_apuracao(regime_apuracao) == "caixa":
        cursor.execute(
            """
            SELECT MIN(TO_CHAR(data_pagamento, 'YYYY-MM')) AS c
            FROM financeiro_mensalidades
            WHERE LOWER(COALESCE(status, '')) = 'pago' AND data_pagamento IS NOT NULL
            """
        )
    else:
        cursor.execute(
            """
            SELECT MIN(TO_CHAR(data_vencimento, 'YYYY-MM')) AS c
            FROM financeiro_mensalidades
            WHERE LOWER(COALESCE(status, '')) NOT IN ('cancelado', 'cancelada')
            """
        )
    row = cursor.fetchone() or {}
    return row.get("c")


def receita_sistema_mes(cursor, competencia, regime_apuracao="competencia"):
    if normalizar_regime_apuracao(regime_apuracao) == "caixa":
        cursor.execute(
            """
            SELECT COALESCE(SUM(
                valor::numeric + COALESCE(juros_valor, 0) + COALESCE(multa_valor, 0)
            ), 0) AS total
            FROM financeiro_mensalidades
            WHERE LOWER(COALESCE(status, '')) = 'pago'
              AND data_pagamento IS NOT NULL
              AND TO_CHAR(data_pagamento, 'YYYY-MM') = %s
            """,
            (competencia,),
        )
    else:
        cursor.execute(
            """
            SELECT COALESCE(SUM(valor::numeric), 0) AS total
            FROM financeiro_mensalidades
            WHERE TO_CHAR(data_vencimento, 'YYYY-MM') = %s
              AND LOWER(COALESCE(status, '')) NOT IN ('cancelado', 'cancelada')
            """,
            (competencia,),
        )
    return float((cursor.fetchone() or {}).get("total") or 0)


def resumo_emitido_e_caixa(cursor, competencia):
    cursor.execute(
        """
        SELECT
            COALESCE(SUM(CASE
                WHEN TO_CHAR(data_vencimento, 'YYYY-MM') = %s
                 AND LOWER(COALESCE(status, '')) NOT IN ('cancelado', 'cancelada')
                THEN valor::numeric ELSE 0 END), 0) AS emitido,
            COALESCE(SUM(CASE
                WHEN LOWER(COALESCE(status, '')) = 'pago'
                 AND data_pagamento IS NOT NULL
                 AND TO_CHAR(data_pagamento, 'YYYY-MM') = %s
                THEN valor::numeric + COALESCE(juros_valor, 0) + COALESCE(multa_valor, 0) ELSE 0 END), 0) AS recebido_caixa,
            COALESCE(SUM(CASE
                WHEN TO_CHAR(data_vencimento, 'YYYY-MM') = %s
                 AND LOWER(COALESCE(status, '')) NOT IN ('pago', 'cancelado', 'cancelada')
                THEN valor::numeric ELSE 0 END), 0) AS inadimplente
        FROM financeiro_mensalidades
        """,
        (competencia, competencia, competencia),
    )
    row = cursor.fetchone() or {}
    return {
        "emitido": float(row.get("emitido") or 0),
        "recebido_caixa": float(row.get("recebido_caixa") or 0),
        "inadimplente": float(row.get("inadimplente") or 0),
    }


def listar_recebimentos_mes(cursor, competencia):
    cursor.execute(
        """
        SELECT f.data_pagamento, f.data_vencimento, f.descricao, f.forma_pagamento, f.valor,
               f.parcela_contrato, f.juros_percentual, f.juros_valor, f.multa_valor,
               a.matricula, a.nome_completo
        FROM financeiro_mensalidades f
        LEFT JOIN alunos a ON a.id = f.aluno_id
        WHERE LOWER(COALESCE(f.status, '')) = 'pago'
          AND f.data_pagamento IS NOT NULL
          AND TO_CHAR(f.data_pagamento, 'YYYY-MM') = %s
        ORDER BY f.data_pagamento, a.nome_completo
        """,
        (competencia,),
    )
    return cursor.fetchall() or []


def folha_sistema_mes(cursor, competencia):
    cursor.execute(
        """
        SELECT COALESCE(SUM(custo_escola::numeric), 0) AS total
        FROM folha_itens
        WHERE competencia = %s
        """,
        (competencia,),
    )
    return float((cursor.fetchone() or {}).get("total") or 0)


def listar_folha_janela(cursor, competencias):
    competencias = [item for item in (competencias or []) if item]
    if not competencias:
        return []
    cursor.execute(
        """
        SELECT fi.competencia, fi.tipo_contrato, fi.bruto, fi.encargos, fi.custo_escola,
               fi.liquido, fi.inss_funcionario, fi.irrf, f.nome_completo
        FROM folha_itens fi
        LEFT JOIN funcionarios f ON f.id = fi.funcionario_id
        WHERE fi.competencia = ANY(%s)
        ORDER BY fi.competencia, f.nome_completo
        """,
        (competencias,),
    )
    return cursor.fetchall() or []


def mapa_competencias(cursor):
    cursor.execute(
        """
        SELECT competencia, receita_bruta, folha_encargos, origem, observacao
        FROM simples_competencias
        ORDER BY competencia
        """
    )
    return {row["competencia"]: dict(row) for row in (cursor.fetchall() or [])}


def _receita_da_linha(cursor, comp, gravado, regime_apuracao):
    if gravado and (gravado.get("origem") or "") in _ORIGENS_CONGELADAS:
        return float(gravado.get("receita_bruta") or 0)
    return receita_sistema_mes(cursor, comp, regime_apuracao)


def montar_quadro_simples(cursor, mes_apuracao, regime_apuracao="competencia"):
    regime_apuracao = normalizar_regime_apuracao(regime_apuracao)
    janela = janela_competencias(mes_apuracao)
    primeira = primeira_competencia_sistema(cursor, regime_apuracao)
    gravados = mapa_competencias(cursor)
    linhas = []
    rbt12 = 0.0
    fs12 = 0.0
    meses_validos = 0
    for comp in janela:
        fora_primeira = bool(primeira and comp < primeira)
        gravado = gravados.get(comp) or {}
        rec = _receita_da_linha(cursor, comp, gravado, regime_apuracao)
        folha = float(gravado.get("folha_encargos") or 0) if gravado else folha_sistema_mes(cursor, comp)
        entra = not fora_primeira
        if entra:
            rbt12 += rec
            fs12 += folha
            meses_validos += 1
        linhas.append(
            {
                "competencia": comp,
                "receita_bruta": rec,
                "folha_encargos": folha,
                "origem": gravado.get("origem") or ("sistema" if rec or folha else ""),
                "observacao": gravado.get("observacao") or "",
                "entra_rbt12": entra,
                "rotulo": _rotulo_comp(comp),
            }
        )
    gravado_mes = gravados.get(mes_apuracao) or {}
    receita_mes = _receita_da_linha(cursor, mes_apuracao, gravado_mes, regime_apuracao)
    folha_mes = float(gravado_mes.get("folha_encargos") or 0) or folha_sistema_mes(cursor, mes_apuracao)
    linha_apuracao = {
        "competencia": mes_apuracao,
        "receita_bruta": receita_mes,
        "folha_encargos": folha_mes,
        "origem": gravado_mes.get("origem") or ("sistema" if receita_mes or folha_mes else ""),
        "observacao": gravado_mes.get("observacao") or "",
        "entra_rbt12": False,
        "rotulo": _rotulo_comp(mes_apuracao),
    }
    return {
        "janela": janela,
        "linhas": linhas,
        "linha_apuracao": linha_apuracao,
        "rbt12": rbt12,
        "fs12": fs12,
        "meses_validos": meses_validos or 1,
        "primeira": primeira,
        "receita_mes": receita_mes,
        "folha_mes": folha_mes,
        "mes_apuracao": mes_apuracao,
        "regime_apuracao": regime_apuracao,
    }


def _rotulo_comp(comp):
    meses = (
        "jan", "fev", "mar", "abr", "mai", "jun",
        "jul", "ago", "set", "out", "nov", "dez",
    )
    try:
        ano, mes = parse_mes(comp)
        return f"{meses[mes - 1]}/{ano}"
    except Exception:
        return comp


def carregar_sistema(cursor, mes_apuracao, regime_apuracao="competencia"):
    comps = janela_competencias(mes_apuracao) + [mes_apuracao]
    qtd = 0
    for comp in comps:
        rec = receita_sistema_mes(cursor, comp, regime_apuracao)
        folha = folha_sistema_mes(cursor, comp)
        if rec or folha:
            upsert_competencia(cursor, comp, rec, folha, "sistema", "Carga do sistema (mensalidades e folha)")
            qtd += 1
    return qtd


def carregar_folhas(cursor, mes_apuracao):
    comps = janela_competencias(mes_apuracao) + [mes_apuracao]
    qtd = 0
    for comp in comps:
        folha = folha_sistema_mes(cursor, comp)
        if folha:
            upsert_competencia(cursor, comp, None, folha, "folha", "Folha e encargos do sistema")
            qtd += 1
    return qtd


def _decodificar_csv(bruto):
    if bruto.startswith(b"\xff\xfe") or bruto.startswith(b"\xfe\xff"):
        return bruto.decode("utf-16")
    if bruto.startswith(b"\xef\xbb\xbf"):
        return bruto.decode("utf-8-sig")
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            return bruto.decode(enc)
        except UnicodeDecodeError:
            continue
    return bruto.decode("latin-1", errors="replace")


def _separador_csv(texto):
    linhas = [ln for ln in texto.splitlines() if ln.strip()][:8]
    if not linhas:
        return ";"
    pontos = sum(ln.count(";") for ln in linhas)
    virgulas = sum(ln.count(",") for ln in linhas)
    tabs = sum(ln.count("\t") for ln in linhas)
    if tabs > pontos and tabs > virgulas:
        return "\t"
    if pontos >= virgulas and pontos > 0:
        return ";"
    if virgulas > 0:
        return ","
    return ";"


def _linhas_arquivo(arquivo):
    nome = (arquivo.filename or "").lower()
    bruto = arquivo.read()
    if nome.endswith(".xlsx") or nome.endswith(".xls"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ValueError("Para planilha Excel instale openpyxl, ou salve o arquivo como CSV.")
        wb = load_workbook(io.BytesIO(bruto), data_only=True)
        ws = wb.active
        return [[cell if cell is not None else "" for cell in row] for row in ws.iter_rows(values_only=True)]
    texto = _decodificar_csv(bruto)
    return list(csv.reader(io.StringIO(texto), delimiter=_separador_csv(texto)))


def importar_competencias(arquivo, origem="planilha"):
    linhas = _linhas_arquivo(arquivo)
    if not linhas:
        return []
    cab = [str(c or "").strip().lower() for c in linhas[0]]
    idx_comp = next((i for i, c in enumerate(cab) if any(k in c for k in ("compet", "periodo", "período", "mes", "mês", "apurac"))), None)
    idx_rec = next((i for i, c in enumerate(cab) if any(k in c for k in ("receita", "rbt", "fatur", "valor", "bruta"))), None)
    idx_folha = next((i for i, c in enumerate(cab) if any(k in c for k in ("folha", "encargo", "fs12", "massa"))), None)
    tem_cabecalho = idx_comp is not None
    saida = []
    for row in linhas[1 if tem_cabecalho else 0:]:
        if not row or not any(str(c).strip() for c in row if c is not None):
            continue
        if tem_cabecalho:
            comp = parse_competencia(row[idx_comp] if idx_comp < len(row) else "")
            rec = parse_moeda_livre(row[idx_rec] if idx_rec is not None and idx_rec < len(row) else 0)
            folha = parse_moeda_livre(row[idx_folha] if idx_folha is not None and idx_folha < len(row) else 0)
        else:
            comp = None
            rec = 0.0
            folha = 0.0
            for cell in row:
                if comp is None:
                    talvez = parse_competencia(cell)
                    if talvez:
                        comp = talvez
                        continue
                if rec <= 0:
                    v = parse_moeda_livre(cell)
                    if v > 0:
                        rec = v
            if len(row) >= 3:
                folha = parse_moeda_livre(row[2])
        if comp:
            saida.append({"competencia": comp, "receita_bruta": rec, "folha_encargos": folha, "origem": origem})
    return saida


def gravar_importacao(cursor, itens):
    for item in itens:
        upsert_competencia(
            cursor,
            item["competencia"],
            item.get("receita_bruta"),
            item.get("folha_encargos"),
            item.get("origem") or "planilha",
            "Importado",
        )
    return len(itens)
